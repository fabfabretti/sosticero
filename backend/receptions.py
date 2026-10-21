import base64
import binascii
from datetime import date
from difflib import SequenceMatcher
from io import BytesIO
import re
import unicodedata
from zipfile import ZipFile, BadZipFile
from xml.etree.ElementTree import ParseError
from defusedxml.common import DefusedXmlException

from .domain import DAYS, HOURS, SITES


def empty():
    return dict(schema=1, versions=[])


def normalize(value):
    value = unicodedata.normalize('NFKD', str(value)).casefold()
    return re.sub(r'[^a-z0-9]', '', value)


def on_day(document, day):
    selected = next((v for v in reversed(document['versions']) if v['effective'] <= day), None)
    return selected['rows'] if selected else []


def at(ds, teacher, day, hour):
    weekday = DAYS[date.fromisoformat(day).weekday()]
    return next((r for r in ds.get('receptions', []) if r['teacher'] == teacher and r['day'] == weekday and r['hour'] == hour), None)


def replace(document, effective, rows):
    if date.fromisoformat(effective).isoformat() != effective:
        raise ValueError('Decorrenza non valida.')
    versions = [v for v in document['versions'] if v['effective'] != effective]
    return dict(schema=1, versions=sorted(versions + [dict(effective=effective, rows=rows)], key=lambda v: v['effective']))


def validate(document, catalog):
    if not isinstance(document, dict) or set(document) != {'schema', 'versions'} or type(document['schema']) is not int or document['schema'] != 1 or not isinstance(document['versions'], list):
        raise ValueError('Archivio ricevimenti non valido.')
    previous = ''
    for version in document['versions']:
        if not isinstance(version, dict) or set(version) != {'effective', 'rows'}:
            raise ValueError('Versione ricevimenti non valida.')
        effective = version['effective']
        if not isinstance(effective, str) or date.fromisoformat(effective).isoformat() != effective or effective <= previous or not isinstance(version['rows'], list):
            raise ValueError('Decorrenze ricevimenti non valide.')
        previous = effective
        seen = set()
        for row in version['rows']:
            if not isinstance(row, dict) or set(row) != {'teacher', 'day', 'hour', 'name', 'time', 'location', 'site'}:
                raise ValueError('Ricevimento non valido.')
            if not isinstance(row['teacher'], str) or row['teacher'] not in catalog or row['day'] not in DAYS[:5] or type(row['hour']) is not int or row['hour'] not in HOURS or row['site'] not in (*SITES, None):
                raise ValueError('Docente, giorno, ora o sede ricevimento non validi.')
            if any(not isinstance(row[k], str) or len(row[k]) > 500 for k in ('name', 'time', 'location')) or not row['name'].strip():
                raise ValueError('Dettagli ricevimento non validi.')
            key = (row['teacher'], row['day'], row['hour'])
            if key in seen:
                raise ValueError('Ricevimento duplicato.')
            seen.add(key)
    return document


def parse(encoded):
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException
    if not isinstance(encoded, str) or len(encoded) > 8_000_000:
        raise ValueError('File XLSX troppo grande (massimo 6 MB).')
    try:
        content = base64.b64decode(encoded, validate=True)
        with ZipFile(BytesIO(content)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 30_000_000:
                raise ValueError('File XLSX decompresso troppo grande.')
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
    except (ValueError, binascii.Error, BadZipFile, OSError, KeyError, ParseError, DefusedXmlException, InvalidFileException) as exc:
        raise ValueError('File XLSX non valido o troppo grande.') from exc
    try:
        for sheet in workbook:
            values = sheet.iter_rows(values_only=True)
            header = [normalize(v or '') for v in next(values, ())]
            if not {'docente', 'giorno', 'ora'} <= set(header):
                continue
            rows = []
            for number, cells in enumerate(values, 2):
                if number > 2001:
                    raise ValueError('Massimo 2000 righe ricevimenti.')
                if not any(v is not None for v in cells):
                    continue
                fields = dict(zip(header, cells))
                name = str(fields.get('docente') or '').strip()
                day = normalize(fields.get('giorno') or '')[:3]
                raw_hour = fields.get('ora')
                hour = int(raw_hour) if not isinstance(raw_hour, bool) and str(raw_hour).isdigit() else None
                error = '' if name and len(name) <= 500 and day in DAYS[:5] and hour in HOURS else 'Docente, giorno o ora non validi.'
                location = str(fields.get('luogo') or '').strip()
                time = str(fields.get('orario') or '').strip()
                if len(location) > 500 or len(time) > 500:
                    error = 'Dettagli troppo lunghi.'
                site = {'centrale': 'centrale', 'sedecentrale': 'centrale', 'succursale': 'succursale'}.get(normalize(location))
                rows.append(dict(row=number, name=name, day=day, hour=hour, time=time, location=location, site=site, error=error))
            return rows
        raise ValueError('Colonne richieste: docente, giorno, ora.')
    except (BadZipFile, KeyError, TypeError, ParseError, DefusedXmlException) as exc:
        raise ValueError('File XLSX non valido.') from exc
    finally:
        workbook.close()


def review(rows, teachers, names, choices):
    if not isinstance(choices, dict):
        raise ValueError('Abbinamenti non validi.')
    aliases = {identifier: {normalize(t['name']), *(normalize(v) for v in names['teachers'].get(identifier, {}).values())}
               for identifier, t in teachers.items() if not t.get('retired')}
    if any(key not in {str(r['row']) for r in rows} for key in choices):
        raise ValueError('Riga abbinamento non valida.')
    result, saved, seen = [], [], {}
    for source in rows:
        row = dict(source)
        wanted = normalize(row['name'])
        exact = [identifier for identifier, labels in aliases.items() if wanted and wanted in labels]
        ranked = sorted(aliases, key=lambda identifier: max(SequenceMatcher(None, wanted, label).ratio() for label in aliases[identifier]), reverse=True)
        choice = choices.get(str(row['row']), {})
        if not isinstance(choice, dict) or set(choice) - {'teacher', 'site', 'skip'} or ('skip' in choice and type(choice['skip']) is not bool):
            raise ValueError('Scelta ricevimento non valida.')
        teacher = choice.get('teacher', exact[0] if len(exact) == 1 else None)
        site = choice.get('site', row['site'])
        if teacher is not None and (not isinstance(teacher, str) or teacher not in aliases) or site not in (*SITES, None):
            raise ValueError('Docente o sede ricevimento non validi.')
        row.update(teacher=teacher, site=site, skip=choice.get('skip', False), suggestions=ranked[:5])
        if not row['skip'] and not row['error'] and teacher:
            entry = {k: row[k] for k in ('teacher', 'day', 'hour', 'name', 'time', 'location', 'site')}
            key = (teacher, row['day'], row['hour'])
            if key in seen and seen[key] != entry:
                row['error'] = 'Ricevimenti in conflitto nella stessa ora.'
            elif key not in seen:
                seen[key] = entry
                saved.append(entry)
        result.append(row)
    ready = all(row['skip'] or row['teacher'] and not row['error'] for row in result)
    return result, saved, bool(ready)
